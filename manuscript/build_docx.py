"""Build manuscript_inline.docx (inline figures + table after first citation,
Times New Roman 12pt, double-spaced, APA refs) for Religion submission."""
import os, re
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from latex2mathml.converter import convert as latex_to_mathml
from docx_equation import mathml_to_omml

A = os.path.dirname(os.path.abspath(__file__))
R = os.path.dirname(A)
MATH_TEX = r"H_r = \int_T \int_A N(x,t)\, E_r(x,t)\, dx\, dt"
MATH_LIT = "H\u1d63 = \u222b\u209c\u222b_A N(x,t)\u00b7E\u1d63(x,t) dx dt"

FIG = {
    "Figure 1": (f"{R}/figures/fig1_release_cycle.png",
                 "Figure 1. The normative release cycle."),
    "Figure 2": (f"{R}/figures/fig2_persistence_models.png",
                 "Figure 2. Four accounts of persistence after religious recession, and normative release."),
    "Figure 3": (f"{R}/figures/fig3_revival.png",
                 "Figure 3. Revival without restoration."),
}
TABLE_MD = []  # populated below


def emit_runs(p, text):
    """Markdown emphasis -> runs; math literal -> OMML."""
    text = text.replace("H\u1d63 = \u222b\u209c\u222b_A N(x,t)\u00b7E\u1d63(x,t) dx dt",
                        "\u0000MATH\u0000")
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|\u0000MATH\u0000)", text):
        if not tok:
            continue
        if tok == "\u0000MATH\u0000":
            p._element.append(mathml_to_omml(latex_to_mathml(MATH_TEX)))
        elif tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            p.add_run(tok)


def add_para(doc, text, style=None, size=12, bold=False, center=False,
             double=True, italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(6)
    if double:
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit_runs(p, text)
    for r in p.runs:
        r.font.name = "Times New Roman"; r.font.size = Pt(size)
        if bold: r.bold = True
        if italic: r.italic = True
    return p


def add_fig(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(5.8))
    c = add_para(doc, caption, size=10, double=False)
    return


def add_table(doc, rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = tbl.cell(i, j).paragraphs[0]
            emit_runs(c, cell)
            for r in c.runs:
                r.font.name = "Times New Roman"; r.font.size = Pt(9)
                if i == 0 or j == 0 and i == 0:
                    r.font.bold = (i == 0)


def main():
    md = open(f"{A}/manuscript.md").read()
    # parse the Table 1 markdown block
    t1 = re.search(r"(\| Framework \|.+?)(?=\n## )", md, re.S).group(1)
    rows = [[c.strip() for c in l.strip().strip("|").split("|")]
            for l in t1.strip().splitlines() if not re.match(r"^\|[\s\-|]+\|$", l)]
    # split sections: body before "## Acknowledgements", then refs, then table/captions
    body = md.split("## Acknowledgements")[0]
    refs = md.split("## References")[1].split("## Tables")[0]

    doc = docx.Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)

    # title
    lines = body.splitlines()
    add_para(doc, lines[0].lstrip("# "), size=14, bold=True, center=True)
    add_para(doc, "[Author name and affiliation omitted]", center=True, italic=True)

    inserted = set()
    table_done = False
    for block in re.split(r"\n\s*\n", body.split("\n", 1)[1]):
        b = block.strip()
        if not b:
            continue
        if b.startswith("## "):
            add_para(doc, b[3:], size=13, bold=True, double=False)
        elif b.startswith("**Keywords:**"):
            add_para(doc, b)
        elif b.startswith("|"):
            continue  # handled later
        elif b.startswith("**Figure"):
            continue
        else:
            # bullet-ish lines
            if b.startswith("*"):
                for item in b.splitlines():
                    item = item.strip()
                    if item.startswith("*") and not item.startswith("**"):
                        p = add_para(doc, item)
                        p.paragraph_format.left_indent = Inches(0.4)
                        continue
                    add_para(doc, item)
            else:
                add_para(doc, b)
        # insert figures after first citation paragraph
        for name, (path, cap) in FIG.items():
            if name not in inserted and name in b:
                inserted.add(name)
                add_fig(doc, path, cap)
        if not table_done and ("(Figure 2; Table 1)" in b or "Table 1" in b):
            table_done = True
            add_para(doc, "Table 1. Accounts of persistence after religious recession: "
                          "what persists, what recedes, and how persistence is interpreted.",
                     size=10, double=False)
            add_table(doc, rows)

    # back matter
    add_para(doc, "Acknowledgements", size=13, bold=True, double=False)
    add_para(doc, "[To be completed.]")
    add_para(doc, "Disclosure statement", size=13, bold=True, double=False)
    add_para(doc, "No potential conflict of interest was reported by the author.")
    add_para(doc, "References", size=13, bold=True, double=False)
    for ent in re.split(r"\n\s*\n", refs.strip()):
        e = " ".join(l.strip() for l in ent.splitlines())
        p = doc.add_paragraph()
        emit_runs(p, e)
        for r in p.runs:
            r.font.name = "Times New Roman"; r.font.size = Pt(12)
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.left_indent = Inches(0.5); pf.first_line_indent = Inches(-0.5)

    out = f"{A}/manuscript_inline.docx"
    doc.save(out)
    print("saved", out, "figs inserted:", sorted(inserted))


main()
